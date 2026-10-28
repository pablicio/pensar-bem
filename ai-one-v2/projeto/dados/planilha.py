"""A planilha da Marlene: RATEIO_FINAL_v7_AGORA_VAI.xlsx."""
from datetime import datetime

from openpyxl import Workbook

from cadastro import zip_fixo
from ged import cpf_enviado, creditos, geracao
from mundo import BRUTOS, rng

NOMES_USINA = {"corumba": "Corumbá de Goiás",
               "cocalzinho": "Cocalzinho",
               "abadiania": "Abadiânia"}


def gravar_planilha(pessoas, energia, presente):
    # o rateio de fevereiro como a planilha calculou
    fevereiro = (337, 365)
    usada = geracao(energia, presente, *fevereiro,
                    planilha=True)
    enviado = creditos(pessoas, usada)
    wb = Workbook()
    par = wb.active
    par.title = "Parâmetros"
    par.append(["Usina", "Geração do mês (kWh)"])
    for usina, nome in NOMES_USINA.items():
        par.append([nome, usada[usina]])
    par.append([])
    par.append(["Tarifa (R$/kWh)", 0.9412])
    visivel = wb.create_sheet("Rateio")
    escondida = wb.create_sheet("NÃO APAGAR")
    escondida.sheet_state = "hidden"
    cabecalho = ["CPF", "Nome", "Usina", "Cota (%)", "kWh",
                 "R$"]
    visivel.append(cabecalho)
    escondida.append(cabecalho)
    coladas = set(int(i) for i in
                  rng.choice(len(pessoas), 14, replace=False))
    for i, p in enumerate(pessoas):
        n = i + 2
        usina = NOMES_USINA[p["usina"]]
        busca = f"VLOOKUP(C{n},Parâmetros!$A$2:$B$4,2,0)"
        kwh = f"=ROUND({busca}*D{n}/100,2)"
        reais = f"=ROUND(E{n}*Parâmetros!$B$6,2)"
        escondida.append([cpf_enviado(p), p["nome"], usina,
                          p["cota"], kwh, reais])
        if i in coladas:             # valor de novembro colado
            fator = rng.uniform(0.8, 1.2)
            kwh = round(enviado[p["cpf"]][0] * fator, 2)
            reais = round(kwh * 0.9412, 2)
        visivel.append([cpf_enviado(p), p["nome"], usina,
                        p["cota"], kwh, reais])
    wb.properties.creator = "Marlene"
    wb.properties.created = datetime(2026, 9, 14, 17, 52)
    arquivo = BRUTOS / "RATEIO_FINAL_v7_AGORA_VAI.xlsx"
    wb.save(arquivo)
    zip_fixo(arquivo)
