from openpyxl import load_workbook

GERACAO = 126809.824     # kWh de Abadiânia em fevereiro
TARIFA = 0.9412          # R$ por kWh

arquivo = "dados/brutos/cadastro_cooperados.xlsx"
aba = load_workbook(arquivo)["Cooperados"]
cooperados = list(aba.iter_rows(values_only=True))[1:]

abadiania = [l for l in cooperados
             if l[4] == "Abadiânia" and l[7] == "Ativo"]
print(len(abadiania), "cooperados ativos em Abadiânia")

total = 0
por_tipo = {"residencial": 0, "comercio": 0, "rural": 0}
for i, linha in enumerate(abadiania, start=1):
    nome, tipo, cota = linha[1], linha[3], linha[5]
    kwh = round(GERACAO * cota / 100, 2)
    total += kwh
    por_tipo[tipo] += kwh
    if i <= 5:
        reais = round(kwh * TARIFA, 2)
        print(f"{i}. {nome:26} {cota:6.4f}% {kwh:7.2f} kWh "
              f"R$ {reais:7.2f}")

for tipo, kwh in por_tipo.items():
    print(f"{tipo:12} {kwh:10.2f} kWh")
print(f"soma dos créditos: {total:.2f} kWh")
print(f"geração do mês:    {GERACAO:.2f} kWh")
print(f"sobra:             {GERACAO - total:.3f} kWh")
