from openpyxl import load_workbook

arquivo = "dados/brutos/cadastro_cooperados.xlsx"
aba = load_workbook(arquivo)["Cooperados"]
cooperados = list(aba.iter_rows(values_only=True))[1:]

cotas = sorted([l[5] for l in cooperados
                if l[4] == "Abadiânia" and l[7] == "Ativo"],
               reverse=True)
soma = 0
n = 0
while soma < 50:
    soma += cotas[n]
    n += 1
print(n, "de", len(cotas), "cooperados somam",
      round(soma, 2), "%")
