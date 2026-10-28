from openpyxl import load_workbook

arquivo = "dados/brutos/cadastro_cooperados.xlsx"
aba = load_workbook(arquivo)["Cooperados"]
cooperados = list(aba.iter_rows(values_only=True))[1:]

cotas = [l[5] for l in cooperados
         if l[4] == "Abadiânia" and l[7] == "Ativo"
         and l[3] == "comercio"]
soma = sum(cotas)
print(len(cotas), "comércios")
print(f"cota somada: {soma:.4f}%")
print(f"crédito: {126809.824 * soma / 100:.2f} kWh")
