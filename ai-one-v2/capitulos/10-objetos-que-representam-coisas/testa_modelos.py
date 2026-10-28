from openpyxl import load_workbook

from modelos import Cooperado, Leitura

arquivo = "dados/brutos/cadastro_cooperados.xlsx"
aba = load_workbook(arquivo)["Cooperados"]
dirceus = []
for linha in list(aba.iter_rows(values_only=True))[1:]:
    if linha[1] == "Dirceu Alves" and linha[4] == "Abadiânia":
        dirceus.append(Cooperado(cpf=linha[2], nome=linha[1],
                                 tipo=linha[3], usina=linha[4],
                                 cota_pct=linha[5],
                                 situacao=linha[7]))

antigo, novo = dirceus
print(antigo.cpf, antigo.situacao, antigo.cota_pct)
print(novo.cpf, novo.situacao, novo.cota_pct)
print("são o mesmo cooperado?", antigo == novo)
print("crédito de fevereiro:", novo.credito(126809.824))

meio_dia = Leitura("Corumbá de Goiás", "INV-10",
                   "10/02/2027 12:00", 1.0473)
print(meio_dia.plausivel(potencia_kw=39.0625))
