from openpyxl import load_workbook

from modelos import Cooperado

arquivo = "dados/brutos/cadastro_cooperados.xlsx"
aba = load_workbook(arquivo)["Cooperados"]
por_cpf = {}
pares = 0
for linha in list(aba.iter_rows(values_only=True))[1:]:
    c = Cooperado(linha[2], linha[1], linha[3], linha[4],
                  linha[5], linha[7])
    if c.cpf in por_cpf and por_cpf[c.cpf] == c:
        pares += 1
    por_cpf[c.cpf] = c
print(pares, "pares de linhas são o mesmo cooperado")
