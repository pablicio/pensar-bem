from openpyxl import load_workbook

from funcoes import normaliza_cpf

arquivo = "dados/brutos/cadastro_cooperados.xlsx"
aba = load_workbook(arquivo)["Cooperados"]
cooperados = list(aba.iter_rows(values_only=True))[1:]

vezes = {}
for linha in cooperados:
    cpf = normaliza_cpf(linha[2])
    vezes[cpf] = vezes.get(cpf, 0) + 1
repetidos = [cpf for cpf, n in vezes.items() if n > 1]
print(len(cooperados), "linhas")
print(len(vezes), "CPFs diferentes")
print(len(repetidos), "CPFs em mais de uma linha")
