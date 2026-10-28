from openpyxl import load_workbook

arquivo = "dados/brutos/RATEIO_FINAL_v7_AGORA_VAI.xlsx"
planilha = load_workbook(arquivo)
print(planilha.sheetnames)
print(planilha["NÃO APAGAR"].sheet_state)

visivel = planilha["Rateio"].iter_rows(min_row=2,
                                         values_only=True)
escondida = planilha["NÃO APAGAR"].iter_rows(min_row=2,
                                             values_only=True)
diferentes = 0
for n, (v, e) in enumerate(zip(visivel, escondida), start=2):
    if v[4] != e[4]:
        diferentes += 1
        if diferentes <= 2:
            print(n, v[1], "| na frente:", v[4])
            print("   atrás:", e[4])
print(diferentes, "linhas diferentes")
