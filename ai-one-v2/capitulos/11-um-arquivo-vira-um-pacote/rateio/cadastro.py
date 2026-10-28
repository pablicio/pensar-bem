"""O cadastro de cooperados, da planilha para objetos."""
from openpyxl import load_workbook

from rateio.modelos import Cooperado


def ler_cooperados(caminho):
    """Todas as linhas da aba Cooperados, como Cooperado."""
    aba = load_workbook(caminho)["Cooperados"]
    linhas = list(aba.iter_rows(values_only=True))[1:]
    return [Cooperado(cpf=l[2], nome=l[1], tipo=l[3],
                      usina=l[4], cota_pct=l[5], situacao=l[7])
            for l in linhas]
