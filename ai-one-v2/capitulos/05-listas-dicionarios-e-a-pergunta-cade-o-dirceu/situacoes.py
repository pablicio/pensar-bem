from openpyxl import load_workbook

arquivo = "dados/brutos/cadastro_cooperados.xlsx"
aba = load_workbook(arquivo)["Cooperados"]
cooperados = list(aba.iter_rows(values_only=True))[1:]

contagem = {}
for linha in cooperados:
    situacao = linha[7]
    contagem[situacao] = contagem.get(situacao, 0) + 1
print(contagem)
