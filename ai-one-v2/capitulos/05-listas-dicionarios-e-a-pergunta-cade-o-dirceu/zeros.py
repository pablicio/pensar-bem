from openpyxl import load_workbook

arquivo = "dados/brutos/cadastro_cooperados.xlsx"
aba = load_workbook(arquivo)["Cooperados"]
cooperados = list(aba.iter_rows(values_only=True))[1:]

com_zero = 0
como_numero = 0
for linha in cooperados:
    if linha[7] != "Ativo":
        continue
    cpf = linha[2]
    if not isinstance(cpf, str):
        como_numero += 1
    digitos = str(cpf).strip().replace(".", "")
    digitos = digitos.replace("-", "")
    if digitos.zfill(11).startswith("0"):
        com_zero += 1
print("ativos com CPF começando em zero:", com_zero)
print("ativos com CPF guardado como número:", como_numero)
