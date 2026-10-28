from openpyxl import load_workbook

arquivo = "dados/brutos/cadastro_cooperados.xlsx"
aba = load_workbook(arquivo)["Cooperados"]
linhas = list(aba.iter_rows(values_only=True))

cabecalho = linhas[0]
cooperados = linhas[1:]
print(len(cooperados), "linhas no cadastro")
print(cabecalho[:5])
print(cooperados[0][:3])

for linha in cooperados:
    if linha[1] == "Dirceu Alves":
        print(linha[0], linha[2], linha[4], linha[7])
