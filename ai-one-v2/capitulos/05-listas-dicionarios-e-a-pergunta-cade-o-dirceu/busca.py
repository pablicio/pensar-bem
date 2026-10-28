from openpyxl import load_workbook

arquivo = "dados/brutos/cadastro_cooperados.xlsx"
aba = load_workbook(arquivo)["Cooperados"]
cooperados = list(aba.iter_rows(values_only=True))[1:]

por_cpf = {}
for linha in cooperados:
    por_cpf[linha[2]] = linha
print(len(por_cpf), "CPFs diferentes no dicionário")

documento = "060.419.765-98"      # da carteira do Seu Dirceu
print(por_cpf[documento][1], por_cpf[documento][7])
print("06041976598" in por_cpf)
print(6041976598 in por_cpf)
